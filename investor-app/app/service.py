"""Logika bisnis: membaca/menulis basis data dan memanggil mesin perhitungan."""
from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import threading
import urllib.request
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal as D, InvalidOperation
from pathlib import Path

from . import calc
from .calc import Biaya, GalatHitung, NAMA_JENIS, q0, q2
from .db import RAIZ, path_db

ZONA = timezone(timedelta(hours=int(os.environ.get("INVESTOR_UTC_OFFSET", "7"))))
PENGATURAN_AWAL = {
    "persen_senin": "4", "persen_selasa": "4", "persen_rabu": "4",
    "persen_kamis": "4", "persen_jumat": "4", "persen_sabtu": "2",
    "persen_minggu": "0",
    "biaya_penarikan_persen": "6", "biaya_konversi_persen": "0.5",
    "biaya_transfer_idr": "11000",
}
KURS_AWAL = "16000"
_kunci = threading.RLock()
_cache: dict[str, dict] = {}


class Galat(Exception):
    """Kesalahan yang pesannya boleh ditampilkan ke pengguna."""

    def __init__(self, pesan: str, kode: int = 400):
        super().__init__(pesan)
        self.pesan = pesan
        self.kode = kode


# --------------------------------------------------------------------------
# Pembantu umum
# --------------------------------------------------------------------------
def sekarang() -> datetime:
    return datetime.now(ZONA)


def hari_ini() -> date:
    return sekarang().date()


def _d(teks) -> D:
    return D(str(teks))


def ambil_desimal(data: dict, kunci: str, label: str, *, minimum=None, maksimum=None,
                  lebih_dari=None, wajib=True, bawaan=None) -> D | None:
    mentah = data.get(kunci)
    if mentah is None or (isinstance(mentah, str) and mentah.strip() == ""):
        if wajib:
            raise Galat(f"{label} harus diisi.")
        return bawaan
    if isinstance(mentah, bool):
        raise Galat(f"{label} harus berupa angka.")
    try:
        nilai = D(str(mentah).strip())
        if not nilai.is_finite():
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        raise Galat(f"{label} harus berupa angka.")
    if lebih_dari is not None and nilai <= lebih_dari:
        raise Galat(f"{label} harus lebih besar dari {lebih_dari}.")
    if minimum is not None and nilai < minimum:
        raise Galat(f"{label} tidak boleh kurang dari {minimum}.")
    if maksimum is not None and nilai > maksimum:
        raise Galat(f"{label} tidak boleh lebih dari {maksimum}.")
    return nilai


def ambil_tanggal(data: dict, kunci: str, label: str, *, wajib=True, bawaan=None) -> date | None:
    mentah = data.get(kunci)
    if mentah is None or (isinstance(mentah, str) and not mentah.strip()):
        if wajib:
            raise Galat(f"{label} harus diisi.")
        return bawaan
    try:
        return date.fromisoformat(str(mentah).strip()[:10])
    except ValueError:
        raise Galat(f"{label} tidak valid. Gunakan format tahun-bulan-tanggal, contoh 2026-10-07.")


def ambil_teks(data: dict, kunci: str, label: str, *, wajib=False, maks=500) -> str:
    teks = str(data.get(kunci) or "").strip()
    if wajib and not teks:
        raise Galat(f"{label} harus diisi.")
    if len(teks) > maks:
        raise Galat(f"{label} terlalu panjang (maksimal {maks} huruf).")
    return teks


def _batas_tanggal(tgl: date, label: str, hi: date | None = None):
    hi = hi or hari_ini()
    if tgl > hi:
        raise Galat(f"{label} tidak boleh di masa depan (setelah {calc.fmt_tgl(hi)}).")
    if tgl.year < 2000:
        raise Galat(f"{label} tidak valid.")


# --------------------------------------------------------------------------
# Pengaturan & kurs
# --------------------------------------------------------------------------
def muat_pengaturan(conn) -> dict:
    baris = {r["kunci"]: r["nilai"] for r in conn.execute("SELECT kunci, nilai FROM pengaturan")}
    for k, v in PENGATURAN_AWAL.items():
        baris.setdefault(k, v)
    persen = [_d(baris[f"persen_{h}"]) for h in calc.KUNCI_HARI]
    biaya = Biaya(_d(baris["biaya_penarikan_persen"]), _d(baris["biaya_konversi_persen"]),
                  _d(baris["biaya_transfer_idr"]))
    return {"persen_hari": persen, "biaya": biaya, "kurs": kurs_terakhir(conn)["nilai"]}


def kurs_terakhir(conn) -> dict:
    r = conn.execute("SELECT * FROM kurs_riwayat ORDER BY tanggal DESC, id DESC LIMIT 1").fetchone()
    if not r:
        return {"tanggal": hari_ini(), "jam": "00:00", "nilai": _d(KURS_AWAL), "sumber": "awal"}
    return {"tanggal": date.fromisoformat(r["tanggal"]), "jam": r["jam"],
            "nilai": _d(r["nilai"]), "sumber": r["sumber"], "id": r["id"]}


def _fungsi_kurs(conn):
    rows = [(date.fromisoformat(r["tanggal"]), _d(r["nilai"])) for r in
            conn.execute("SELECT tanggal, nilai FROM kurs_riwayat ORDER BY tanggal, id")]
    if not rows:
        return lambda _t: _d(KURS_AWAL)
    tgls = [r[0] for r in rows]

    def kurs_pada(t: date) -> D:
        # kurs terakhir pada atau sebelum tanggal t; jika belum ada, kurs paling awal
        lo, hi = 0, len(tgls)
        while lo < hi:
            mid = (lo + hi) // 2
            if tgls[mid] <= t:
                lo = mid + 1
            else:
                hi = mid
        return rows[lo - 1][1] if lo > 0 else rows[0][1]
    return kurs_pada


def tampil_pengaturan(conn) -> dict:
    p = muat_pengaturan(conn)
    k = kurs_terakhir(conn)
    return {
        "kurs": k["nilai"], "kurs_tanggal": k["tanggal"], "kurs_jam": k["jam"], "kurs_sumber": k["sumber"],
        "persen": {h: p["persen_hari"][i] for i, h in enumerate(calc.KUNCI_HARI)},
        "biaya_penarikan_persen": p["biaya"].penarikan_persen,
        "biaya_konversi_persen": p["biaya"].konversi_persen,
        "biaya_transfer_idr": p["biaya"].transfer_idr,
    }


def simpan_pengaturan(conn, data: dict) -> dict:
    persen_baru = {}
    for h in calc.KUNCI_HARI:
        persen_baru[h] = ambil_desimal(
            (data.get("persen") or {}), h, f"Keuntungan hari {h.capitalize()}",
            minimum=D("-100"), maksimum=D("100"))
        if persen_baru[h] <= -100:
            raise Galat(f"Keuntungan hari {h.capitalize()} harus lebih besar dari -100%.")
    bp = ambil_desimal(data, "biaya_penarikan_persen", "Biaya Penarikan", minimum=D(0), maksimum=D(100))
    bk = ambil_desimal(data, "biaya_konversi_persen", "Biaya Konversi", minimum=D(0), maksimum=D(100))
    bt = ambil_desimal(data, "biaya_transfer_idr", "Biaya Transfer", minimum=D(0), maksimum=D(100000000))
    pasangan = {f"persen_{h}": str(v) for h, v in persen_baru.items()}
    pasangan.update({"biaya_penarikan_persen": str(bp), "biaya_konversi_persen": str(bk),
                     "biaya_transfer_idr": str(bt)})
    with conn:
        for k, v in pasangan.items():
            conn.execute("INSERT INTO pengaturan(kunci, nilai) VALUES(?, ?) "
                         "ON CONFLICT(kunci) DO UPDATE SET nilai = excluded.nilai", (k, v))
        kurs_baru = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0),
                                  maksimum=D(10000000), wajib=False)
        if kurs_baru is not None and kurs_baru != kurs_terakhir(conn)["nilai"]:
            _catat_kurs(conn, kurs_baru, "manual")
    batalkan_cache(conn)
    return tampil_pengaturan(conn)


def _catat_kurs(conn, nilai: D, sumber: str):
    n = sekarang()
    conn.execute("INSERT INTO kurs_riwayat(tanggal, jam, nilai, sumber) VALUES(?,?,?,?)",
                 (n.date().isoformat(), n.strftime("%H:%M"), str(nilai), sumber))


def simpan_kurs(conn, data: dict) -> dict:
    nilai = ambil_desimal(data, "nilai", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000))
    with conn:
        _catat_kurs(conn, nilai, "manual")
    batalkan_cache(conn)
    return tampil_pengaturan(conn)


def perbarui_kurs_online(conn) -> dict:
    """Ambil kurs terbaru dari internet. Gagal dengan sopan bila offline."""
    sumber_url = "https://open.er-api.com/v6/latest/USD"
    try:
        req = urllib.request.Request(sumber_url, headers={"User-Agent": "pelacak-modal/1.0"})
        with urllib.request.urlopen(req, timeout=8) as r:
            hasil = json.loads(r.read().decode("utf-8"))
        nilai = D(str(hasil["rates"]["IDR"]))
        if nilai <= 0:
            raise ValueError("kurs tidak valid")
    except Exception as e:  # noqa: BLE001 - jaringan bisa gagal dengan banyak cara
        return {"ok": False, "pesan": "Kurs terbaru tidak dapat diambil (internet tidak tersedia atau "
                                      "sumber tidak merespons). Kurs manual tetap dipakai.",
                "detail": type(e).__name__}
    nilai = nilai.quantize(D("0.01"))
    with conn:
        _catat_kurs(conn, nilai, "open.er-api.com")
    batalkan_cache(conn)
    return {"ok": True, "pesan": f"Kurs diperbarui menjadi Rp{nilai:,.2f} per USD.", **tampil_pengaturan(conn)}


def riwayat_kurs(conn, batas=50) -> list[dict]:
    return [{"id": r["id"], "tanggal": r["tanggal"], "jam": r["jam"], "nilai": _d(r["nilai"]),
             "sumber": r["sumber"]} for r in
            conn.execute("SELECT * FROM kurs_riwayat ORDER BY tanggal DESC, id DESC LIMIT ?", (batas,))]


# --------------------------------------------------------------------------
# Transaksi: baca & hitung
# --------------------------------------------------------------------------
def _trx_dict(r) -> dict:
    return {
        "id": r["id"], "investor_id": r["investor_id"], "tanggal": date.fromisoformat(r["tanggal"]),
        "jenis": r["jenis"], "jumlah_usd": _d(r["jumlah_usd"]), "kurs": _d(r["kurs"]),
        "nilai_idr": _d(r["nilai_idr"]), "biaya_penarikan_usd": _d(r["biaya_penarikan_usd"]),
        "biaya_konversi_usd": _d(r["biaya_konversi_usd"]), "nilai_bersih_usd": _d(r["nilai_bersih_usd"]),
        "biaya_transfer_idr": _d(r["biaya_transfer_idr"]),
        "bersih_idr_diterima": _d(r["bersih_idr_diterima"]), "catatan": r["catatan"],
    }


def muat_transaksi(conn, investor_id: int) -> list[dict]:
    return [_trx_dict(r) for r in conn.execute(
        "SELECT * FROM transaksi WHERE investor_id = ? ORDER BY tanggal, id", (investor_id,))]


def muat_aktual(conn) -> dict[date, D]:
    return {date.fromisoformat(r["tanggal"]): _d(r["persen"])
            for r in conn.execute("SELECT tanggal, persen FROM hasil_aktual")}


def hitung_investor(conn, investor_id: int, tgl_akhir: date | None = None, *,
                    pakai_aktual=True, tambahan_trx: list[dict] | None = None,
                    tanpa_id: int | None = None) -> list[dict]:
    p = muat_pengaturan(conn)
    trx = [t for t in muat_transaksi(conn, investor_id) if t["id"] != tanpa_id]
    if tambahan_trx:
        trx = trx + tambahan_trx
    return calc.bangun_garis_waktu(trx, p["persen_hari"], muat_aktual(conn) if pakai_aktual else None,
                                   _fungsi_kurs(conn), p["biaya"], tgl_akhir or hari_ini(), pakai_aktual)


def batalkan_cache(conn=None):
    with _kunci:
        _cache.pop(path_db(), None)


def semua_garis_waktu(conn, *, pakai_aktual=True) -> dict[int, list[dict]]:
    """Garis waktu semua investor. Disimpan sementara di memori dan dibuat ulang
    setiap ada perubahan data atau pergantian hari."""
    kunci = path_db()
    with _kunci:
        c = _cache.get(kunci)
        tgl = hari_ini()
        if c and c["tanggal"] == tgl and pakai_aktual in c["data"]:
            return c["data"][pakai_aktual]
        if not c or c["tanggal"] != tgl:
            c = {"tanggal": tgl, "data": {}}
        hasil = {}
        p = muat_pengaturan(conn)
        aktual = muat_aktual(conn) if pakai_aktual else None
        kurs_fn = _fungsi_kurs(conn)
        for inv in conn.execute("SELECT id FROM investor ORDER BY id"):
            trx = muat_transaksi(conn, inv["id"])
            hasil[inv["id"]] = calc.bangun_garis_waktu(
                trx, p["persen_hari"], aktual, kurs_fn, p["biaya"], tgl, pakai_aktual)
        c["data"][pakai_aktual] = hasil
        _cache[kunci] = c
        if pakai_aktual:
            _simpan_snapshot(conn, hasil, "perhitungan ulang")
        return hasil


KOLOM_SNAPSHOT = ["saldo_awal", "tambahan", "penyesuaian", "penarikan", "saldo_dasar", "persen",
                  "sumber_persen", "keuntungan", "saldo_akhir", "kurs", "saldo_idr", "total_modal_usd",
                  "total_keuntungan_usd", "total_penarikan_usd", "kekayaan_bersih_idr", "persen_untung"]


def _simpan_snapshot(conn, garis: dict[int, list[dict]], alasan: str):
    """Simpan hasil perhitungan harian ke basis data + catat riwayat perhitungan."""
    masalah = []
    jumlah = 0
    with conn:
        conn.execute("DELETE FROM perhitungan_harian")
        for iid, baris in garis.items():
            masalah += [f"investor {iid}: {m}" for m in calc.periksa_identitas(baris)]
            conn.executemany(
                f"INSERT INTO perhitungan_harian(investor_id, tanggal, {', '.join(KOLOM_SNAPSHOT)}) "
                f"VALUES (?, ?, {', '.join('?' * len(KOLOM_SNAPSHOT))})",
                [(iid, b["tanggal"].isoformat(), *[str(b[k]) for k in KOLOM_SNAPSHOT]) for b in baris])
            jumlah += len(baris)
        conn.execute(
            "INSERT INTO riwayat_perhitungan(waktu, alasan, jumlah_investor, jumlah_baris, status, pesan) "
            "VALUES (?,?,?,?,?,?)",
            (sekarang().isoformat(timespec="seconds"), alasan, len(garis), jumlah,
             "ok" if not masalah else "ada masalah", "; ".join(masalah[:5])))
        conn.execute("DELETE FROM riwayat_perhitungan WHERE id NOT IN "
                     "(SELECT id FROM riwayat_perhitungan ORDER BY id DESC LIMIT 300)")


def periksa_integritas(conn) -> dict:
    """Pemeriksaan internal: hitung ulang dari transaksi, bandingkan dengan
    hasil yang tersimpan, dan pastikan saldo hanya berubah karena transaksi
    atau keuntungan harian yang tercatat."""
    masalah: list[str] = []
    tersimpan = {}
    for r in conn.execute("SELECT investor_id, tanggal, saldo_akhir FROM perhitungan_harian"):
        tersimpan[(r["investor_id"], r["tanggal"])] = _d(r["saldo_akhir"])
    jumlah_baris = 0
    for inv in conn.execute("SELECT id, kode FROM investor ORDER BY id"):
        try:
            baris = hitung_investor(conn, inv["id"])
        except GalatHitung as e:
            masalah.append(f"Investor {inv['kode']}: {e}")
            continue
        jumlah_baris += len(baris)
        masalah += [f"Investor {inv['kode']}: {m}" for m in calc.periksa_identitas(baris)]
        trx = muat_transaksi(conn, inv["id"])
        if baris:
            setoran = sum((t["jumlah_usd"] for t in trx if t["jenis"] != "penarikan"), D(0))
            tarik = sum((t["jumlah_usd"] for t in trx if t["jenis"] == "penarikan"), D(0))
            untung = sum((b["keuntungan"] for b in baris), D(0))
            if setoran - tarik + untung != baris[-1]["saldo_akhir"]:
                masalah.append(f"Investor {inv['kode']}: saldo akhir tidak sama dengan "
                               f"setoran - penarikan + keuntungan tercatat")
        for b in baris:
            k = (inv["id"], b["tanggal"].isoformat())
            if k not in tersimpan:
                masalah.append(f"Investor {inv['kode']}: perhitungan {b['tanggal']} belum tersimpan")
            elif tersimpan[k] != b["saldo_akhir"]:
                masalah.append(f"Investor {inv['kode']}: saldo tersimpan {b['tanggal']} berbeda dari hasil hitung ulang")
    return {"ok": not masalah, "masalah": masalah[:50], "jumlah_masalah": len(masalah),
            "jumlah_baris": jumlah_baris, "diperiksa_pada": sekarang().isoformat(timespec="seconds")}


# --------------------------------------------------------------------------
# Ringkasan investor
# --------------------------------------------------------------------------
def _status_warna(ring: dict) -> str:
    if ring["total_modal_idr"] <= 0:
        return "abu"
    if ring["titik_impas"]["tercapai"]:
        return "hijau"
    if ring["kekayaan_idr"] >= ring["total_modal_idr"] * D("0.8"):
        return "kuning"
    return "merah"


def ringkasan(conn, inv, baris: list[dict], p: dict | None = None) -> dict:
    p = p or muat_pengaturan(conn)
    kurs = p["kurs"]
    hi = hari_ini()
    trx = muat_transaksi(conn, inv["id"])
    modal_awal = sum((t["jumlah_usd"] for t in trx if t["jenis"] == "modal_awal"), D(0))
    tambahan = sum((t["jumlah_usd"] for t in trx if t["jenis"] == "tambahan"), D(0))
    penyesuaian = sum((t["jumlah_usd"] for t in trx if t["jenis"] == "penyesuaian"), D(0))
    a = baris[-1] if baris else None
    saldo = a["saldo_akhir"] if a else D(0)
    # nilai saat ini memakai kurs terbaru, bukan kurs hari terakhir di tabel
    rincian = calc.hitung_penarikan(saldo, kurs, p["biaya"]) if saldo > 0 else calc.hitung_penarikan(0, kurs, p["biaya"])
    cair = rincian["bersih_idr_diterima"]
    tarik_bersih = a["total_penarikan_bersih_idr"] if a else D(0)
    modal_idr = a["total_modal_idr"] if a else D(0)
    kekayaan = tarik_bersih + cair
    untung_bersih = kekayaan - modal_idr
    persen = (untung_bersih / modal_idr * 100) if modal_idr > 0 else D(0)
    if baris:
        # pakai nilai terkini (kurs terbaru) pada hari terakhir untuk status
        akhir = dict(a)
        akhir["nilai_bersih_cair_idr"] = cair
        akhir["kekayaan_bersih_idr"] = kekayaan
        baris_status = baris[:-1] + [akhir]
    else:
        baris_status = baris
    target = calc.ringkas_target(baris_status, p["persen_hari"], p["biaya"], kurs, hi)
    r = {
        "id": inv["id"], "kode": inv["kode"], "nama": inv["nama"],
        "tanggal_mulai": date.fromisoformat(inv["tanggal_mulai"]),
        "status": inv["status"], "catatan": inv["catatan"],
        "modal_awal_usd": modal_awal, "tambahan_usd": tambahan, "penyesuaian_usd": penyesuaian,
        "total_modal_usd": a["total_modal_usd"] if a else D(0), "total_modal_idr": modal_idr,
        "total_penarikan_usd": a["total_penarikan_usd"] if a else D(0),
        "uang_diterima_idr": tarik_bersih, "uang_diterima_usd": q2(tarik_bersih / kurs),
        "saldo_usd": saldo, "saldo_idr": q0(saldo * kurs),
        "total_keuntungan_usd": a["total_keuntungan_usd"] if a else D(0),
        "nilai_total_usd": saldo + (a["total_penarikan_usd"] if a else D(0)),
        "nilai_bersih_cair_idr": cair, "nilai_bersih_cair_usd": q2(cair / kurs),
        "kekayaan_idr": kekayaan, "kekayaan_usd": q2(kekayaan / kurs),
        "untung_bersih_idr": untung_bersih, "untung_bersih_usd": q2(untung_bersih / kurs),
        "persen_untung": persen,
        "rincian_cair": rincian,
        "titik_impas": target[0],
        "untung_100": target[100],
        "target": [target[t] for t in calc.TARGET_PERSEN if t != 0],
        "hari_berjalan": (hi - date.fromisoformat(inv["tanggal_mulai"])).days + 1,
    }
    r["warna"] = _status_warna(r)
    return r


def _semua_investor(conn):
    return conn.execute("SELECT * FROM investor ORDER BY kode, id").fetchall()


def dashboard(conn) -> dict:
    garis = semua_garis_waktu(conn)
    p = muat_pengaturan(conn)
    kurs = p["kurs"]
    daftar = [ringkasan(conn, inv, garis[inv["id"]], p) for inv in _semua_investor(conn)]
    s = lambda k: sum((x[k] for x in daftar), D(0))  # noqa: E731
    modal_idr = s("total_modal_idr")
    kekayaan = s("kekayaan_idr")
    tot = {
        "jumlah_investor": len(daftar),
        "jumlah_aktif": sum(1 for x in daftar if x["status"] == "aktif"),
        "total_modal_usd": s("total_modal_usd"), "total_modal_idr": modal_idr,
        "total_saldo_usd": s("saldo_usd"), "total_saldo_idr": s("saldo_idr"),
        "total_keuntungan_usd": s("total_keuntungan_usd"),
        "total_penarikan_usd": s("total_penarikan_usd"),
        "uang_diterima_idr": s("uang_diterima_idr"), "uang_diterima_usd": q2(s("uang_diterima_idr") / kurs),
        "nilai_bersih_cair_idr": s("nilai_bersih_cair_idr"),
        "nilai_bersih_cair_usd": q2(s("nilai_bersih_cair_idr") / kurs),
        "kekayaan_idr": kekayaan, "kekayaan_usd": q2(kekayaan / kurs),
        "untung_bersih_idr": kekayaan - modal_idr, "untung_bersih_usd": q2((kekayaan - modal_idr) / kurs),
        "persen_untung": ((kekayaan - modal_idr) / modal_idr * 100) if modal_idr > 0 else D(0),
        "jumlah_titik_impas": sum(1 for x in daftar if x["titik_impas"]["tercapai"]),
        "jumlah_untung_100": sum(1 for x in daftar if x["untung_100"]["tercapai"]),
    }
    return {"tanggal": hari_ini(), "kurs": tampil_pengaturan(conn), "total": tot, "investor": daftar}


# --------------------------------------------------------------------------
# Investor
# --------------------------------------------------------------------------
def _ambil_investor(conn, investor_id) -> sqlite3.Row:
    try:
        iid = int(investor_id)
    except (TypeError, ValueError):
        raise Galat("Pilih investor terlebih dahulu.")
    r = conn.execute("SELECT * FROM investor WHERE id = ?", (iid,)).fetchone()
    if not r:
        raise Galat("Investor tidak ditemukan.", 404)
    return r


def kode_berikutnya(conn) -> str:
    kode_ada = {r["kode"] for r in conn.execute("SELECT kode FROM investor")}
    for huruf in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
        if huruf not in kode_ada:
            return huruf
    n = 1
    while f"INV-{n:03d}" in kode_ada:
        n += 1
    return f"INV-{n:03d}"


def _validasi_investor(conn, data, *, id_sendiri=None):
    kode = ambil_teks(data, "kode", "Kode Investor", wajib=True, maks=20)
    nama = ambil_teks(data, "nama", "Nama Investor", wajib=True, maks=100)
    status = str(data.get("status") or "aktif")
    if status not in ("aktif", "selesai"):
        raise Galat("Status harus Aktif atau Selesai.")
    catatan = ambil_teks(data, "catatan", "Catatan", maks=1000)
    dobel = conn.execute("SELECT id FROM investor WHERE lower(kode) = lower(?)", (kode,)).fetchone()
    if dobel and dobel["id"] != id_sendiri:
        raise Galat(f"Kode Investor '{kode}' sudah dipakai. Gunakan kode lain.")
    return kode, nama, status, catatan


def tambah_investor(conn, data: dict) -> dict:
    data = dict(data)
    if not str(data.get("kode") or "").strip():
        data["kode"] = kode_berikutnya(conn)
    kode, nama, status, catatan = _validasi_investor(conn, data)
    mulai = ambil_tanggal(data, "tanggal_mulai", "Tanggal Mulai")
    _batas_tanggal(mulai, "Tanggal Mulai")
    modal = ambil_desimal(data, "modal_awal_usd", "Modal Awal USD", lebih_dari=D(0), maksimum=D("1000000000"))
    kurs = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000),
                         wajib=False, bawaan=muat_pengaturan(conn)["kurs"])
    modal = q2(modal)
    with conn:
        cur = conn.execute("INSERT INTO investor(kode, nama, tanggal_mulai, status, catatan, dibuat) "
                           "VALUES (?,?,?,?,?,?)",
                           (kode, nama, mulai.isoformat(), status, catatan, sekarang().isoformat(timespec="seconds")))
        iid = cur.lastrowid
        _sisip_trx(conn, iid, mulai, "modal_awal", modal, kurs, "Modal awal")
        hitung_investor(conn, iid)  # validasi
    batalkan_cache(conn)
    return detail_investor(conn, iid)


def ubah_investor(conn, investor_id, data: dict) -> dict:
    inv = _ambil_investor(conn, investor_id)
    kode, nama, status, catatan = _validasi_investor(conn, data, id_sendiri=inv["id"])
    mulai = ambil_tanggal(data, "tanggal_mulai", "Tanggal Mulai", wajib=False,
                          bawaan=date.fromisoformat(inv["tanggal_mulai"]))
    _batas_tanggal(mulai, "Tanggal Mulai")
    with conn:
        conn.execute("UPDATE investor SET kode=?, nama=?, status=?, catatan=? WHERE id=?",
                     (kode, nama, status, catatan, inv["id"]))
        awal = conn.execute("SELECT * FROM transaksi WHERE investor_id=? AND jenis='modal_awal'",
                            (inv["id"],)).fetchone()
        modal_baru = ambil_desimal(data, "modal_awal_usd", "Modal Awal USD", lebih_dari=D(0),
                                   maksimum=D("1000000000"), wajib=False)
        if awal and (mulai.isoformat() != awal["tanggal"] or modal_baru is not None):
            jumlah = q2(modal_baru) if modal_baru is not None else _d(awal["jumlah_usd"])
            conn.execute("UPDATE transaksi SET tanggal=?, jumlah_usd=?, nilai_idr=? WHERE id=?",
                         (mulai.isoformat(), str(jumlah), str(q0(jumlah * _d(awal["kurs"]))), awal["id"]))
            conn.execute("UPDATE investor SET tanggal_mulai=? WHERE id=?", (mulai.isoformat(), inv["id"]))
        _pastikan_urutan(conn, inv["id"])
        hitung_investor(conn, inv["id"])
    batalkan_cache(conn)
    return detail_investor(conn, inv["id"])


def hapus_investor(conn, investor_id) -> dict:
    inv = _ambil_investor(conn, investor_id)
    with conn:
        conn.execute("DELETE FROM investor WHERE id = ?", (inv["id"],))
    batalkan_cache(conn)
    return {"ok": True, "pesan": f"Investor {inv['nama']} beserta seluruh transaksinya sudah dihapus."}


def daftar_investor(conn) -> list[dict]:
    return dashboard(conn)["investor"]


def detail_investor(conn, investor_id) -> dict:
    inv = _ambil_investor(conn, investor_id)
    garis = semua_garis_waktu(conn)
    r = ringkasan(conn, inv, garis[inv["id"]])
    r["transaksi"] = [_tampil_trx(t) for t in reversed(muat_transaksi(conn, inv["id"]))]
    r["kurs"] = tampil_pengaturan(conn)
    return r


# --------------------------------------------------------------------------
# Transaksi: tulis
# --------------------------------------------------------------------------
def _tampil_trx(t: dict, kode: str | None = None, nama: str | None = None) -> dict:
    o = dict(t)
    o["jenis_nama"] = NAMA_JENIS[t["jenis"]]
    if kode is not None:
        o["investor_kode"], o["investor_nama"] = kode, nama
    return o


def _sisip_trx(conn, iid, tgl: date, jenis, jumlah: D, kurs: D, catatan, rincian=None):
    rincian = rincian or {}
    cur = conn.execute(
        "INSERT INTO transaksi(investor_id, tanggal, jenis, jumlah_usd, kurs, nilai_idr, biaya_penarikan_usd, "
        "biaya_konversi_usd, nilai_bersih_usd, biaya_transfer_idr, bersih_idr_diterima, catatan, dibuat) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (iid, tgl.isoformat(), jenis, str(jumlah), str(kurs), str(q0(jumlah * kurs)),
         str(rincian.get("biaya_penarikan_usd", 0)), str(rincian.get("biaya_konversi_usd", 0)),
         str(rincian.get("nilai_bersih_usd", 0)), str(rincian.get("biaya_transfer_idr", 0)),
         str(rincian.get("bersih_idr_diterima", 0)), catatan, sekarang().isoformat(timespec="seconds")))
    return cur.lastrowid


def _pastikan_urutan(conn, iid):
    """Semua transaksi harus pada atau setelah Tanggal Mulai (tanggal Modal Awal)."""
    inv = conn.execute("SELECT tanggal_mulai FROM investor WHERE id=?", (iid,)).fetchone()
    awal = conn.execute("SELECT tanggal FROM transaksi WHERE investor_id=? AND jenis='modal_awal'", (iid,)).fetchall()
    if len(awal) != 1:
        raise Galat("Setiap investor harus memiliki tepat satu Modal Awal.")
    if awal[0]["tanggal"] != inv["tanggal_mulai"]:
        conn.execute("UPDATE investor SET tanggal_mulai=? WHERE id=?", (awal[0]["tanggal"], iid))
    lebih_awal = conn.execute("SELECT 1 FROM transaksi WHERE investor_id=? AND tanggal < ?",
                              (iid, awal[0]["tanggal"])).fetchone()
    if lebih_awal:
        raise Galat("Ada transaksi sebelum Tanggal Mulai investor. Tanggal transaksi harus pada atau "
                    "setelah Tanggal Mulai.")


def _baca_trx_modal(conn, data: dict):
    inv = _ambil_investor(conn, data.get("investor_id"))
    jenis = str(data.get("jenis") or "")
    if jenis not in calc.JENIS_MODAL:
        raise Galat("Pilih jenis transaksi: Modal Awal, Tambahan Modal, atau Penyesuaian.")
    tgl = ambil_tanggal(data, "tanggal", "Tanggal")
    _batas_tanggal(tgl, "Tanggal transaksi")
    if jenis == "penyesuaian":
        jumlah = ambil_desimal(data, "jumlah_usd", "Jumlah USD", minimum=D("-1000000000"), maksimum=D("1000000000"))
        if jumlah == 0:
            raise Galat("Jumlah penyesuaian tidak boleh nol.")
    else:
        jumlah = ambil_desimal(data, "jumlah_usd", "Jumlah USD", lebih_dari=D(0), maksimum=D("1000000000"))
    kurs = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000),
                         wajib=False, bawaan=muat_pengaturan(conn)["kurs"])
    return inv, jenis, tgl, q2(jumlah), kurs, ambil_teks(data, "catatan", "Catatan")


def tambah_modal(conn, data: dict) -> dict:
    inv, jenis, tgl, jumlah, kurs, catatan = _baca_trx_modal(conn, data)
    mulai = date.fromisoformat(inv["tanggal_mulai"])
    if jenis == "modal_awal":
        raise Galat("Investor ini sudah memiliki Modal Awal. Gunakan 'Tambahan Modal' untuk menambah modal, "
                    "atau ubah Modal Awal lewat daftar transaksi.")
    if tgl < mulai:
        raise Galat(f"Tanggal transaksi tidak boleh sebelum Tanggal Mulai investor ({calc.fmt_tgl(mulai)}).")
    with conn:
        tid = _sisip_trx(conn, inv["id"], tgl, jenis, jumlah, kurs, catatan)
        hitung_investor(conn, inv["id"])
    batalkan_cache(conn)
    return {"ok": True, "id": tid, "pesan": "Transaksi modal tersimpan. Saldo investor sudah diperbarui."}


def _cek_penarikan(conn, data: dict, tanpa_id=None) -> dict:
    inv = _ambil_investor(conn, data.get("investor_id"))
    tgl = ambil_tanggal(data, "tanggal", "Tanggal")
    _batas_tanggal(tgl, "Tanggal penarikan")
    mulai = date.fromisoformat(inv["tanggal_mulai"])
    if tgl < mulai:
        raise Galat(f"Tanggal penarikan tidak boleh sebelum Tanggal Mulai investor ({calc.fmt_tgl(mulai)}).")
    jumlah = q2(ambil_desimal(data, "jumlah_usd", "Jumlah Penarikan USD", lebih_dari=D(0), maksimum=D("1000000000")))
    p = muat_pengaturan(conn)
    kurs = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000),
                         wajib=False, bawaan=p["kurs"])
    rinci = calc.hitung_penarikan(jumlah, kurs, p["biaya"])
    # saldo yang tersedia pada tanggal tsb = saldo awal + setoran - penarikan lain pada hari itu
    baris = hitung_investor(conn, inv["id"], tanpa_id=tanpa_id)
    tersedia = D(0)
    for b in baris:
        if b["tanggal"] == tgl:
            tersedia = b["saldo_dasar"]
            break
    else:
        # belum ada transaksi pada hari itu: saldo akhir hari sebelumnya
        sebelum = [b for b in baris if b["tanggal"] < tgl]
        tersedia = sebelum[-1]["saldo_akhir"] if sebelum else D(0)
    return {"inv": inv, "tanggal": tgl, "jumlah": jumlah, "kurs": kurs, "rincian": rinci,
            "tersedia": tersedia, "catatan": ambil_teks(data, "catatan", "Catatan"), "p": p}


def pratinjau_penarikan(conn, data: dict, tanpa_id=None) -> dict:
    c = _cek_penarikan(conn, data, tanpa_id)
    melebihi = c["jumlah"] > c["tersedia"]
    bisa = not melebihi
    saldo_sesudah = None
    pesan = None
    if melebihi:
        pesan = (f"Penarikan USD {c['jumlah']:,.2f} melebihi saldo tersedia USD {c['tersedia']:,.2f} "
                 f"pada {calc.fmt_tgl(c['tanggal'])}. (Saldo tersedia = saldo awal hari itu + setoran hari itu; "
                 f"keuntungan hari berjalan baru masuk di akhir hari.)")
    else:
        try:
            tambahan = {"tanggal": c["tanggal"], "jenis": "penarikan", "jumlah_usd": c["jumlah"],
                        "nilai_idr": D(0), "bersih_idr_diterima": c["rincian"]["bersih_idr_diterima"]}
            baris = hitung_investor(conn, c["inv"]["id"], tambahan_trx=[tambahan], tanpa_id=tanpa_id)
            saldo_sesudah = baris[-1]["saldo_akhir"]
        except GalatHitung as e:
            bisa = False
            pesan = f"Penarikan ini membuat saldo di hari berikutnya negatif. {e}"
    if bisa and not c["rincian"]["cukup_untuk_transfer"]:
        bisa = False
        pesan = "Jumlah penarikan terlalu kecil: setelah semua biaya, nilai yang diterima tidak positif."
    return {"investor_id": c["inv"]["id"], "investor_nama": c["inv"]["nama"], "tanggal": c["tanggal"],
            "saldo_tersedia_usd": c["tersedia"], "melebihi_saldo": melebihi, "bisa_disimpan": bisa,
            "pesan": pesan, "saldo_sesudah_usd": saldo_sesudah, "rincian": c["rincian"]}


def simpan_penarikan(conn, data: dict, tanpa_id=None) -> dict:
    c = _cek_penarikan(conn, data, tanpa_id)
    r = c["rincian"]
    if c["jumlah"] > c["tersedia"]:
        raise Galat(f"Penarikan USD {c['jumlah']:,.2f} melebihi saldo tersedia USD {c['tersedia']:,.2f} "
                    f"pada {calc.fmt_tgl(c['tanggal'])}. Kurangi jumlah penarikan, atau pilih tanggal penarikan setelah "
                    f"keuntungan hari itu masuk.")
    if not r["cukup_untuk_transfer"]:
        raise Galat("Jumlah penarikan terlalu kecil: setelah semua biaya, nilai yang diterima tidak positif.")
    with conn:
        if tanpa_id:
            conn.execute("DELETE FROM transaksi WHERE id = ?", (tanpa_id,))
        tid = _sisip_trx(conn, c["inv"]["id"], c["tanggal"], "penarikan", c["jumlah"], c["kurs"],
                         c["catatan"], r)
        try:
            hitung_investor(conn, c["inv"]["id"])
        except GalatHitung as e:
            raise Galat(f"Penarikan ini membuat saldo di hari-hari berikutnya menjadi negatif. {e}")
    batalkan_cache(conn)
    return {"ok": True, "id": tid, "rincian": r,
            "pesan": f"Penarikan tersimpan. Nilai bersih Rupiah yang diterima: Rp{r['bersih_idr_diterima']:,.0f}."}


def ubah_transaksi(conn, trx_id, data: dict) -> dict:
    r = conn.execute("SELECT * FROM transaksi WHERE id=?", (trx_id,)).fetchone()
    if not r:
        raise Galat("Transaksi tidak ditemukan.", 404)
    data = dict(data)
    data["investor_id"] = r["investor_id"]
    if r["jenis"] == "penarikan":
        return simpan_penarikan(conn, data, tanpa_id=r["id"])
    data["jenis"] = r["jenis"]
    inv, jenis, tgl, jumlah, kurs, catatan = _baca_trx_modal(conn, data)
    mulai = date.fromisoformat(inv["tanggal_mulai"])
    if jenis != "modal_awal" and tgl < mulai:
        raise Galat(f"Tanggal transaksi tidak boleh sebelum Tanggal Mulai investor ({calc.fmt_tgl(mulai)}).")
    with conn:
        conn.execute("UPDATE transaksi SET tanggal=?, jumlah_usd=?, kurs=?, nilai_idr=?, catatan=? WHERE id=?",
                     (tgl.isoformat(), str(jumlah), str(kurs), str(q0(jumlah * kurs)), catatan, r["id"]))
        _pastikan_urutan(conn, inv["id"])
        try:
            hitung_investor(conn, inv["id"])
        except GalatHitung as e:
            raise Galat(f"Perubahan ini membuat saldo negatif. {e}")
    batalkan_cache(conn)
    return {"ok": True, "pesan": "Transaksi diperbarui. Saldo investor sudah dihitung ulang."}


def hapus_transaksi(conn, trx_id) -> dict:
    r = conn.execute("SELECT * FROM transaksi WHERE id=?", (trx_id,)).fetchone()
    if not r:
        raise Galat("Transaksi tidak ditemukan.", 404)
    if r["jenis"] == "modal_awal":
        raise Galat("Modal Awal tidak dapat dihapus. Ubah nilainya, atau hapus investornya.")
    with conn:
        conn.execute("DELETE FROM transaksi WHERE id=?", (r["id"],))
        try:
            hitung_investor(conn, r["investor_id"])
        except GalatHitung as e:
            raise Galat(f"Transaksi ini tidak dapat dihapus karena saldo setelahnya menjadi negatif. {e}")
    batalkan_cache(conn)
    return {"ok": True, "pesan": "Transaksi dihapus. Saldo investor sudah dihitung ulang."}


def riwayat_transaksi(conn, *, investor_id=None, dari=None, sampai=None, jenis=None, q=None) -> list[dict]:
    sql = ("SELECT t.*, i.kode AS investor_kode, i.nama AS investor_nama FROM transaksi t "
           "JOIN investor i ON i.id = t.investor_id WHERE 1=1")
    par: list = []
    if investor_id:
        sql += " AND t.investor_id = ?"
        par.append(int(investor_id))
    if dari:
        sql += " AND t.tanggal >= ?"
        par.append(dari)
    if sampai:
        sql += " AND t.tanggal <= ?"
        par.append(sampai)
    if jenis:
        sql += " AND t.jenis = ?"
        par.append(jenis)
    if q:
        sql += " AND (lower(t.catatan) LIKE ? OR lower(i.nama) LIKE ? OR lower(i.kode) LIKE ?)"
        par += [f"%{q.lower()}%"] * 3
    sql += " ORDER BY t.tanggal DESC, t.id DESC"
    hasil = []
    for r in conn.execute(sql, par):
        t = _trx_dict(r)
        hasil.append(_tampil_trx(t, r["investor_kode"], r["investor_nama"]))
    return hasil


# --------------------------------------------------------------------------
# Perhitungan harian
# --------------------------------------------------------------------------
def _tampil_baris(b: dict) -> dict:
    o = dict(b)
    return o


def harian_per_tanggal(conn, tgl: date) -> dict:
    garis = semua_garis_waktu(conn)
    p = muat_pengaturan(conn)
    baris = []
    for inv in _semua_investor(conn):
        for b in garis[inv["id"]]:
            if b["tanggal"] == tgl:
                o = _tampil_baris(b)
                o.update(investor_id=inv["id"], investor_kode=inv["kode"], investor_nama=inv["nama"])
                baris.append(o)
                break
    total = {k: sum((b[k] for b in baris), D(0)) for k in
             ("saldo_awal", "tambahan", "penyesuaian", "penarikan", "keuntungan", "saldo_akhir",
              "saldo_idr", "nilai_bersih_cair_idr")}
    return {"tanggal": tgl, "hari": calc.NAMA_HARI[tgl.weekday()], "hari_ini": hari_ini(),
            "persen_asumsi": p["persen_hari"][tgl.weekday()], "baris": baris, "total": total}


def harian_investor(conn, investor_id, dari=None, sampai=None) -> dict:
    inv = _ambil_investor(conn, investor_id)
    baris = semua_garis_waktu(conn)[inv["id"]]
    if dari:
        baris = [b for b in baris if b["tanggal"] >= dari]
    if sampai:
        baris = [b for b in baris if b["tanggal"] <= sampai]
    return {"investor": {"id": inv["id"], "kode": inv["kode"], "nama": inv["nama"]},
            "baris": [_tampil_baris(b) for b in baris]}


# --------------------------------------------------------------------------
# Grafik
# --------------------------------------------------------------------------
def data_grafik(conn, *, investor_ids=None, dari=None, sampai=None, status=None) -> dict:
    garis = semua_garis_waktu(conn)
    invs = [i for i in _semua_investor(conn)
            if (not investor_ids or i["id"] in investor_ids) and (not status or status == "semua" or i["status"] == status)]
    hi = hari_ini()
    awal = min((garis[i["id"]][0]["tanggal"] for i in invs if garis[i["id"]]), default=hi)
    dari = max(dari, awal) if dari else awal
    sampai = min(sampai, hi) if sampai else hi
    sumbu, t = [], dari
    while t <= sampai:
        sumbu.append(t)
        t += timedelta(days=1)
    keluar = []
    for i in invs:
        per = {b["tanggal"]: b for b in garis[i["id"]]}
        deret = {k: [] for k in ("saldo_usd", "saldo_idr", "keuntungan_usd", "modal_usd", "modal_idr",
                                 "kekayaan_usd", "kekayaan_idr", "persen_untung",
                                 "penarikan_kum_usd", "penarikan_hari_usd")}
        for t in sumbu:
            b = per.get(t)
            if b is None:
                # sebelum mulai: kosong (nol untuk jumlah)
                for k in deret:
                    deret[k].append(None if k == "persen_untung" else 0)
                continue
            deret["saldo_usd"].append(b["saldo_akhir"])
            deret["saldo_idr"].append(b["saldo_idr"])
            deret["keuntungan_usd"].append(b["total_keuntungan_usd"])
            deret["modal_usd"].append(b["total_modal_usd"])
            deret["modal_idr"].append(b["total_modal_idr"])
            deret["kekayaan_idr"].append(b["kekayaan_bersih_idr"])
            deret["kekayaan_usd"].append(q2(b["kekayaan_bersih_idr"] / b["kurs"]))
            deret["persen_untung"].append(b["persen_untung"])
            deret["penarikan_kum_usd"].append(b["total_penarikan_usd"])
            deret["penarikan_hari_usd"].append(b["penarikan"])
        keluar.append({"id": i["id"], "kode": i["kode"], "nama": i["nama"], "status": i["status"], **deret})
    return {"tanggal": sumbu, "investor": keluar}


# --------------------------------------------------------------------------
# Hasil aktual
# --------------------------------------------------------------------------
def daftar_aktual(conn) -> list[dict]:
    return [{"tanggal": date.fromisoformat(r["tanggal"]), "persen": _d(r["persen"]), "catatan": r["catatan"]}
            for r in conn.execute("SELECT * FROM hasil_aktual ORDER BY tanggal DESC")]


def simpan_aktual(conn, data: dict) -> dict:
    tgl = ambil_tanggal(data, "tanggal", "Tanggal")
    _batas_tanggal(tgl, "Tanggal hasil aktual")
    persen = ambil_desimal(data, "persen", "Persentase Keuntungan Aktual", minimum=D("-100"), maksimum=D("100"))
    if persen <= -100:
        raise Galat("Persentase Keuntungan Aktual harus lebih besar dari -100%.")
    catatan = ambil_teks(data, "catatan", "Catatan")
    with conn:
        conn.execute("INSERT INTO hasil_aktual(tanggal, persen, catatan) VALUES(?,?,?) "
                     "ON CONFLICT(tanggal) DO UPDATE SET persen=excluded.persen, catatan=excluded.catatan",
                     (tgl.isoformat(), str(persen), catatan))
    batalkan_cache(conn)
    return {"ok": True, "pesan": f"Hasil aktual {calc.fmt_tgl(tgl)} tersimpan. Saldo dihitung ulang."}


def hapus_aktual(conn, tanggal: str) -> dict:
    with conn:
        n = conn.execute("DELETE FROM hasil_aktual WHERE tanggal = ?", (tanggal,)).rowcount
    if not n:
        raise Galat("Data hasil aktual tidak ditemukan.", 404)
    batalkan_cache(conn)
    return {"ok": True, "pesan": "Hasil aktual dihapus. Tanggal itu kembali memakai asumsi."}


def perbandingan_aktual(conn, investor_id=None) -> dict:
    """Saldo berdasarkan asumsi vs berdasarkan aktual (per investor atau gabungan)."""
    ids = [int(investor_id)] if investor_id else [r["id"] for r in conn.execute("SELECT id FROM investor")]
    kurs = muat_pengaturan(conn)["kurs"]
    seri: dict[date, dict] = {}
    for iid in ids:
        _ambil_investor(conn, iid)
        a = hitung_investor(conn, iid, pakai_aktual=True)
        s = hitung_investor(conn, iid, pakai_aktual=False)
        for x, y in zip(a, s):
            e = seri.setdefault(x["tanggal"], {"tanggal": x["tanggal"], "aktual": D(0), "asumsi": D(0),
                                               "kurs": x["kurs"], "persen_aktual_dipakai": x["sumber_persen"] == "aktual"})
            e["aktual"] += x["saldo_akhir"]
            e["asumsi"] += y["saldo_akhir"]
    baris = []
    for t in sorted(seri):
        e = seri[t]
        sel = e["aktual"] - e["asumsi"]
        e["selisih_usd"] = sel
        e["selisih_idr"] = q0(sel * e["kurs"])
        e["selisih_persen"] = (sel / e["asumsi"] * 100) if e["asumsi"] != 0 else D(0)
        baris.append(e)
    akhir = baris[-1] if baris else None
    return {"baris": baris, "ringkas": ({
        "asumsi_usd": akhir["asumsi"], "aktual_usd": akhir["aktual"],
        "selisih_usd": akhir["selisih_usd"], "selisih_idr": q0(akhir["selisih_usd"] * kurs),
        "selisih_persen": akhir["selisih_persen"]} if akhir else None)}


# --------------------------------------------------------------------------
# Simulator, skenario, kalkulator
# --------------------------------------------------------------------------
def _persen_dari_form(data: dict, prefiks="") -> list[D]:
    hasil = []
    for h in calc.KUNCI_HARI:
        hasil.append(ambil_desimal(data, prefiks + h, f"Keuntungan {h.capitalize()}", minimum=D("-99.99"), maksimum=D(100)))
    return hasil


def jalankan_simulator(conn, data: dict) -> dict:
    modal = ambil_desimal(data, "modal_awal", "Modal Awal", lebih_dari=D(0), maksimum=D("1000000000"))
    mulai = ambil_tanggal(data, "tanggal_mulai", "Tanggal Mulai")
    akhir = ambil_tanggal(data, "tanggal_akhir", "Tanggal Akhir Simulasi")
    if akhir < mulai:
        raise Galat("Tanggal Akhir Simulasi tidak boleh sebelum Tanggal Mulai.")
    if (akhir - mulai).days > 3650:
        raise Galat("Rentang simulasi maksimal 10 tahun (3.650 hari).")
    kurs = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000))
    persen = _persen_dari_form(data.get("persen") or {})
    p = muat_pengaturan(conn)
    pakai = bool(data.get("pakai_aktual"))
    h = calc.simulasikan(modal, mulai, akhir, persen, kurs, p["biaya"], muat_aktual(conn), pakai)
    for kunci in ("titik_impas", "untung_100"):
        item = h[kunci]
        if item["tanggal"]:
            item["hari_dari_mulai"] = (item["tanggal"] - mulai).days + 1
    h["target"] = [h["target"][t] for t in calc.TARGET_PERSEN if t != 0]
    h["rincian_cair"] = calc.hitung_penarikan(h["saldo_akhir"], kurs, p["biaya"])
    return h


def analisis_skenario(conn, data: dict) -> dict:
    modal = ambil_desimal(data, "modal_awal", "Modal Awal", lebih_dari=D(0), maksimum=D("1000000000"))
    mulai = ambil_tanggal(data, "tanggal_mulai", "Tanggal Mulai", wajib=False, bawaan=hari_ini())
    kurs = ambil_desimal(data, "kurs", "Kurs USD/Rupiah", lebih_dari=D(0), maksimum=D(10000000),
                         wajib=False, bawaan=muat_pengaturan(conn)["kurs"])
    horizon = [30, 60, 90, 180, 365]
    keluar = []
    for s in data.get("skenario") or []:
        nama = ambil_teks(s, "nama", "Nama Skenario", wajib=True, maks=40)
        hari_kerja = ambil_desimal(s, "senin_jumat", f"Keuntungan Senin–Jumat ({nama})", minimum=D("-99.99"), maksimum=D(100))
        sabtu = ambil_desimal(s, "sabtu", f"Keuntungan Sabtu ({nama})", minimum=D("-99.99"), maksimum=D(100))
        minggu = ambil_desimal(s, "minggu", f"Keuntungan Minggu ({nama})", minimum=D("-99.99"), maksimum=D(100),
                               wajib=False, bawaan=D(0))
        persen = [hari_kerja] * 5 + [sabtu, minggu]
        hasil = calc.hasil_horizon(modal, mulai, persen, horizon)
        p = muat_pengaturan(conn)
        akhir = mulai + timedelta(days=365)
        lengkap = calc.simulasikan(modal, mulai, akhir - timedelta(days=1), persen, kurs, p["biaya"])
        keluar.append({"nama": nama, "persen": {"senin_jumat": hari_kerja, "sabtu": sabtu, "minggu": minggu},
                       "horizon": [{**hasil[h], "saldo_idr": q0(hasil[h]["saldo"] * kurs)} for h in horizon],
                       "grafik": [{"tanggal": b["tanggal"], "saldo": b["saldo_akhir"]} for b in lengkap["baris"]],
                       "titik_impas": lengkap["titik_impas"]})
    if not keluar:
        raise Galat("Tambahkan minimal satu skenario.")
    return {"skenario": keluar, "horizon": horizon}


# --------------------------------------------------------------------------
# Cadangan data
# --------------------------------------------------------------------------
TABEL_CADANGAN = ["investor", "transaksi", "pengaturan", "kurs_riwayat", "hasil_aktual", "riwayat_perhitungan"]


def ekspor_json(conn) -> dict:
    isi = {}
    for t in TABEL_CADANGAN:
        isi[t] = [dict(r) for r in conn.execute(f"SELECT * FROM {t}")]  # noqa: S608 - nama tabel tetap
    return {"aplikasi": "pelacak-modal-investor", "versi": 1,
            "dibuat": sekarang().isoformat(timespec="seconds"), "tabel": isi}


def _csv(kolom, baris) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(kolom)
    for b in baris:
        w.writerow(b)
    return "﻿" + out.getvalue()


def ekspor_csv(conn, tabel: str) -> str:
    if tabel == "investor":
        d = daftar_investor(conn)
        kol = ["Kode", "Nama", "Tanggal Mulai", "Status", "Modal Awal USD", "Tambahan Modal USD",
               "Total Modal USD", "Total Penarikan USD", "Saldo Saat Ini USD", "Total Keuntungan USD",
               "Nilai Bersih Jika Dicairkan Rp", "Kekayaan Bersih Rp", "Persentase Keuntungan",
               "Tanggal Titik Impas", "Tanggal Keuntungan 100%", "Catatan"]
        return _csv(kol, [[x["kode"], x["nama"], x["tanggal_mulai"], x["status"], x["modal_awal_usd"],
                           x["tambahan_usd"], x["total_modal_usd"], x["total_penarikan_usd"], x["saldo_usd"],
                           x["total_keuntungan_usd"], x["nilai_bersih_cair_idr"], x["kekayaan_idr"],
                           round(x["persen_untung"], 2), x["titik_impas"]["tanggal"] or "BELUM TERCAPAI",
                           x["untung_100"]["tanggal"] or "BELUM TERCAPAI", x["catatan"]] for x in d])
    if tabel == "transaksi":
        t = riwayat_transaksi(conn)
        kol = ["Tanggal", "Kode Investor", "Nama Investor", "Jenis", "Jumlah USD", "Kurs", "Nilai Rp",
               "Biaya Penarikan USD", "Biaya Konversi USD", "Nilai Bersih USD", "Biaya Transfer Rp",
               "Bersih Rp Diterima", "Catatan"]
        return _csv(kol, [[x["tanggal"], x["investor_kode"], x["investor_nama"], x["jenis_nama"], x["jumlah_usd"],
                           x["kurs"], x["nilai_idr"], x["biaya_penarikan_usd"], x["biaya_konversi_usd"],
                           x["nilai_bersih_usd"], x["biaya_transfer_idr"], x["bersih_idr_diterima"],
                           x["catatan"]] for x in t])
    if tabel == "harian":
        kol = ["Tanggal", "Hari", "Kode", "Nama Investor", "Saldo Awal", "Tambahan Modal", "Penyesuaian",
               "Penarikan", "Saldo Dasar Perhitungan", "Persentase Keuntungan", "Sumber Persentase",
               "Keuntungan Hari Itu", "Saldo Akhir", "Kurs", "Saldo Rupiah", "Total Modal Disetor",
               "Total Keuntungan", "Total Penarikan", "Kekayaan Bersih Rp", "Persentase Keuntungan Total"]
        garis = semua_garis_waktu(conn)
        baris = []
        for inv in _semua_investor(conn):
            for b in garis[inv["id"]]:
                baris.append([b["tanggal"], b["hari"], inv["kode"], inv["nama"], b["saldo_awal"], b["tambahan"],
                              b["penyesuaian"], b["penarikan"], b["saldo_dasar"], b["persen"], b["sumber_persen"],
                              b["keuntungan"], b["saldo_akhir"], b["kurs"], b["saldo_idr"], b["total_modal_usd"],
                              b["total_keuntungan_usd"], b["total_penarikan_usd"], b["kekayaan_bersih_idr"],
                              round(b["persen_untung"], 4)])
        return _csv(kol, baris)
    raise Galat("Pilih data yang akan diekspor: investor, transaksi, atau harian.")


def ekspor_xlsx(conn) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for judul, kunci in (("Investor", "investor"), ("Transaksi", "transaksi"), ("Perhitungan Harian", "harian")):
        teks = ekspor_csv(conn, kunci).lstrip("﻿")
        ws = wb.create_sheet(judul)
        for i, baris in enumerate(csv.reader(io.StringIO(teks))):
            if i == 0:
                ws.append(baris)
                for c in ws[1]:
                    c.font = Font(bold=True)
            else:
                ws.append([_angka_atau_teks(x) for x in baris])
        ws.freeze_panes = "A2"
        for kol in ws.columns:
            ws.column_dimensions[kol[0].column_letter].width = min(34, max(12, max(len(str(c.value or "")) for c in kol[:50]) + 2))
    ws = wb.create_sheet("Pengaturan")
    ws.append(["Kunci", "Nilai"])
    for r in conn.execute("SELECT kunci, nilai FROM pengaturan ORDER BY kunci"):
        ws.append([r["kunci"], _angka_atau_teks(r["nilai"])])
    ws.append(["kurs_usd_idr", float(kurs_terakhir(conn)["nilai"])])
    ws = wb.create_sheet("Hasil Aktual")
    ws.append(["Tanggal", "Persentase", "Catatan"])
    for a in daftar_aktual(conn):
        ws.append([a["tanggal"].isoformat(), float(a["persen"]), a["catatan"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _angka_atau_teks(x):
    try:
        return float(x) if x not in ("", None) and not str(x)[:1].isalpha() and "-" not in str(x)[1:] else x
    except ValueError:
        return x


def impor_json(conn, isi: dict) -> dict:
    """Pulihkan data dari cadangan JSON. Semua data investor/transaksi/pengaturan
    sekarang diganti. Jika ada yang tidak valid, tidak ada yang berubah."""
    if not isinstance(isi, dict) or isi.get("aplikasi") != "pelacak-modal-investor" or "tabel" not in isi:
        raise Galat("File ini bukan cadangan dari aplikasi ini.")
    tabel = isi["tabel"]
    for t in ("investor", "transaksi", "pengaturan"):
        if t not in tabel or not isinstance(tabel[t], list):
            raise Galat(f"File cadangan tidak lengkap: data '{t}' tidak ditemukan.")
    cadangan_otomatis(conn, "sebelum-impor")
    kolom = {t: [r["name"] for r in conn.execute(f"PRAGMA table_info({t})")] for t in TABEL_CADANGAN}  # noqa: S608
    try:
        with conn:
            conn.execute("PRAGMA defer_foreign_keys = ON")
            for t in ("perhitungan_harian", "transaksi", "investor", "pengaturan", "kurs_riwayat",
                      "hasil_aktual", "riwayat_perhitungan"):
                conn.execute(f"DELETE FROM {t}")  # noqa: S608
            for t in ("investor", "transaksi", "pengaturan", "kurs_riwayat", "hasil_aktual", "riwayat_perhitungan"):
                for baris in tabel.get(t, []):
                    ks = [k for k in baris if k in kolom[t]]
                    if not ks:
                        continue
                    conn.execute(f"INSERT INTO {t}({', '.join(ks)}) VALUES ({', '.join('?' * len(ks))})",  # noqa: S608
                                 [baris[k] for k in ks])
            # validasi: semua investor harus dapat dihitung tanpa saldo negatif
            for inv in conn.execute("SELECT id, kode FROM investor").fetchall():
                try:
                    _pastikan_urutan(conn, inv["id"])
                    hitung_investor(conn, inv["id"])
                except (GalatHitung, Galat) as e:
                    raise Galat(f"Cadangan tidak valid (investor {inv['kode']}): {e}")
    except sqlite3.Error as e:
        raise Galat(f"File cadangan rusak atau tidak sesuai: {e}")
    batalkan_cache(conn)
    semua_garis_waktu(conn)
    return {"ok": True, "pesan": "Data berhasil dipulihkan dari cadangan. Salinan data sebelumnya disimpan otomatis."}


def cadangan_otomatis(conn, alasan="manual") -> str:
    folder = Path(path_db()).parent / "cadangan_otomatis"
    folder.mkdir(parents=True, exist_ok=True)
    nama = folder / f"{sekarang().strftime('%Y%m%d-%H%M%S')}-{alasan}.json"
    nama.write_text(json.dumps(ekspor_json(conn), ensure_ascii=False, indent=1), encoding="utf-8")
    for lama in sorted(folder.glob("*.json"))[:-30]:
        lama.unlink(missing_ok=True)
    return str(nama)


# --------------------------------------------------------------------------
# Data contoh
# --------------------------------------------------------------------------
def isi_data_contoh(conn):
    """Isi data contoh sekali saja, saat basis data masih kosong."""
    if conn.execute("SELECT 1 FROM pengaturan WHERE kunci = '_contoh_terisi'").fetchone():
        return
    with conn:
        conn.execute("INSERT INTO pengaturan(kunci, nilai) VALUES('_contoh_terisi', '1')")
        ada = conn.execute("SELECT 1 FROM investor LIMIT 1").fetchone()
        if ada:
            return
        for k, v in PENGATURAN_AWAL.items():
            conn.execute("INSERT OR IGNORE INTO pengaturan(kunci, nilai) VALUES(?,?)", (k, v))
        if not conn.execute("SELECT 1 FROM kurs_riwayat").fetchone():
            n = sekarang()
            conn.execute("INSERT INTO kurs_riwayat(tanggal, jam, nilai, sumber) VALUES(?,?,?,?)",
                         ((n.date() - timedelta(days=60)).isoformat(), "09:00", KURS_AWAL, "awal"))
        hi = hari_ini()
        kurs = _d(KURS_AWAL)
        contoh = [("A", "Investor A", 40, 1000, "Contoh: modal awal USD 1.000"),
                  ("B", "Investor B", 25, 500, "Contoh: modal awal USD 500"),
                  ("C", "Investor C", 12, 2000, "Contoh: modal awal USD 2.000")]
        ids = {}
        for kode, nama, hari_lalu, modal, cat in contoh:
            mulai = hi - timedelta(days=hari_lalu)
            cur = conn.execute("INSERT INTO investor(kode, nama, tanggal_mulai, status, catatan, dibuat) "
                               "VALUES(?,?,?,?,?,?)", (kode, nama, mulai.isoformat(), "aktif", cat,
                                                       sekarang().isoformat(timespec="seconds")))
            ids[kode] = cur.lastrowid
            _sisip_trx(conn, cur.lastrowid, mulai, "modal_awal", D(modal), kurs, "Modal awal")
        _sisip_trx(conn, ids["B"], hi - timedelta(days=10), "tambahan", D(250), kurs, "Contoh tambahan modal")
        # contoh penarikan Investor A
        p = muat_pengaturan(conn)
        tgl = hi - timedelta(days=5)
        r = calc.hitung_penarikan(D(300), kurs, p["biaya"])
        _sisip_trx(conn, ids["A"], tgl, "penarikan", D(300), kurs, "Contoh penarikan Investor A", r)
    batalkan_cache(conn)
    semua_garis_waktu(conn)

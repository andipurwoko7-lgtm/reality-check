"""Pengujian API: data contoh, validasi, transaksi, cadangan, login."""
import io
import json
from datetime import timedelta

from fastapi.testclient import TestClient

from app import service


def hari(n=0):
    return (service.hari_ini() + timedelta(days=n)).isoformat()


def investor(c, kode):
    return next(i for i in c.get("/api/investor").json()["investor"] if i["kode"] == kode)


def test_login_wajib(tmp_path, monkeypatch):
    monkeypatch.setenv("INVESTOR_DB", str(tmp_path / "x.db"))
    from app.main import app
    with TestClient(app) as c:
        assert c.get("/api/dashboard").status_code == 401
        r = c.post("/api/masuk", json={"nama_pengguna": "admin", "kata_sandi": "salah"})
        assert r.status_code == 401 and "salah" in r.json()["detail"]
        assert c.post("/api/masuk", json={"nama_pengguna": "admin", "kata_sandi": "admin123"}).status_code == 200
        assert c.get("/api/dashboard").status_code == 200
        c.post("/api/keluar")
        assert c.get("/api/dashboard").status_code == 401


def test_kata_sandi_tidak_disimpan_sebagai_teks(klien):
    from app.db import sambung
    conn = sambung()
    h = conn.execute("SELECT kata_sandi_hash FROM pengguna").fetchone()[0]
    assert "admin123" not in h and h.startswith("$2")


def test_batas_percobaan_masuk(tmp_path, monkeypatch):
    monkeypatch.setenv("INVESTOR_DB", str(tmp_path / "x.db"))
    from app.main import app
    with TestClient(app) as c:
        for _ in range(5):
            c.post("/api/masuk", json={"nama_pengguna": "admin", "kata_sandi": "x"})
        assert c.post("/api/masuk", json={"nama_pengguna": "admin", "kata_sandi": "admin123"}).status_code == 429


def test_data_contoh(klien):
    d = klien.get("/api/dashboard").json()
    assert d["total"]["jumlah_investor"] == 3
    kode = {i["kode"]: i for i in d["investor"]}
    assert kode["A"]["modal_awal_usd"] == 1000 and kode["B"]["modal_awal_usd"] == 500 and kode["C"]["modal_awal_usd"] == 2000
    assert kode["B"]["tambahan_usd"] == 250
    assert len({i["tanggal_mulai"] for i in d["investor"]}) == 3
    assert kode["A"]["total_penarikan_usd"] == 300
    assert kode["A"]["uang_diterima_idr"] > 0
    # total konsisten dengan penjumlahan per investor
    assert abs(d["total"]["total_saldo_usd"] - sum(i["saldo_usd"] for i in d["investor"])) < 0.01
    assert klien.get("/api/pemeriksaan").json()["ok"] is True


def test_data_contoh_tidak_diisi_ulang(klien):
    a = investor(klien, "A")
    assert klien.delete(f"/api/investor/{a['id']}").status_code == 200
    from app.main import app
    with TestClient(app) as c2:  # start ulang aplikasi pada DB yang sama
        c2.post("/api/masuk", json={"nama_pengguna": "admin", "kata_sandi": "admin123"})
        assert c2.get("/api/dashboard").json()["total"]["jumlah_investor"] == 2


def test_tambah_ubah_hapus_investor(klien):
    r = klien.post("/api/investor", json={"kode": "D", "nama": "Investor D", "tanggal_mulai": hari(-3),
                                          "modal_awal_usd": 750, "catatan": "uji"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["total_modal_usd"] == 750 and d["saldo_usd"] > 750
    r = klien.put(f"/api/investor/{d['id']}", json={"kode": "D", "nama": "Investor Delta", "status": "selesai",
                                                    "catatan": "x", "modal_awal_usd": 800, "tanggal_mulai": hari(-4)})
    assert r.status_code == 200, r.text
    assert r.json()["nama"] == "Investor Delta" and r.json()["total_modal_usd"] == 800
    assert r.json()["tanggal_mulai"] == hari(-4)
    assert klien.delete(f"/api/investor/{d['id']}").status_code == 200
    assert klien.get(f"/api/investor/{d['id']}").status_code == 404


def test_validasi_investor(klien):
    base = {"kode": "Z", "nama": "Z", "tanggal_mulai": hari(-1), "modal_awal_usd": 100}
    assert "Nama" in klien.post("/api/investor", json={**base, "nama": ""}).json()["detail"]
    assert "Modal Awal" in klien.post("/api/investor", json={**base, "modal_awal_usd": -5}).json()["detail"]
    assert "Modal Awal" in klien.post("/api/investor", json={**base, "modal_awal_usd": 0}).json()["detail"]
    assert "Tanggal" in klien.post("/api/investor", json={**base, "tanggal_mulai": "31-02-2026"}).json()["detail"]
    assert "masa depan" in klien.post("/api/investor", json={**base, "tanggal_mulai": hari(5)}).json()["detail"]
    assert "Kurs" in klien.post("/api/investor", json={**base, "kurs": 0}).json()["detail"]
    assert "sudah dipakai" in klien.post("/api/investor", json={**base, "kode": "a"}).json()["detail"]


def test_transaksi_modal_memperbarui_saldo(klien):
    c = investor(klien, "C")
    r = klien.post("/api/transaksi/modal", json={"investor_id": c["id"], "jenis": "tambahan", "tanggal": hari(),
                                                 "jumlah_usd": 1000, "kurs": 16000})
    assert r.status_code == 200, r.text
    c2 = investor(klien, "C")
    assert c2["total_modal_usd"] == c["total_modal_usd"] + 1000
    assert c2["saldo_usd"] > c["saldo_usd"] + 1000  # + keuntungan hari ini bila Senin-Sabtu atau sama (Minggu)  # noqa


def test_validasi_transaksi_modal(klien):
    c = investor(klien, "C")
    ok = {"investor_id": c["id"], "jenis": "tambahan", "tanggal": hari(), "jumlah_usd": 10, "kurs": 16000}
    assert "investor" in klien.post("/api/transaksi/modal", json={**ok, "investor_id": None}).json()["detail"].lower()
    assert "Jumlah" in klien.post("/api/transaksi/modal", json={**ok, "jumlah_usd": -10}).json()["detail"]
    assert "Kurs" in klien.post("/api/transaksi/modal", json={**ok, "kurs": 0}).json()["detail"]
    assert "Tanggal" in klien.post("/api/transaksi/modal", json={**ok, "tanggal": "bukan-tanggal"}).json()["detail"]
    assert "sebelum Tanggal Mulai" in klien.post("/api/transaksi/modal", json={**ok, "tanggal": hari(-100)}).json()["detail"]
    assert "Modal Awal" in klien.post("/api/transaksi/modal", json={**ok, "jenis": "modal_awal"}).json()["detail"]
    assert "jenis" in klien.post("/api/transaksi/modal", json={**ok, "jenis": "lain"}).json()["detail"].lower()


def test_penarikan_pratinjau_dan_simpan(klien):
    b = investor(klien, "B")
    p = klien.post("/api/penarikan/pratinjau", json={"investor_id": b["id"], "tanggal": hari(), "jumlah_usd": 100,
                                                     "kurs": 16000}).json()
    r = p["rincian"]
    assert r["jumlah_kotor_usd"] == 100 and r["biaya_penarikan_usd"] == 6
    assert r["biaya_konversi_usd"] == 0.47 and r["nilai_bersih_usd"] == 93.53
    assert r["nilai_idr"] == 1496480 and r["bersih_idr_diterima"] == 1485480
    assert p["bisa_disimpan"] and p["saldo_sesudah_usd"] is not None
    # pratinjau tidak mengubah data
    assert investor(klien, "B")["total_penarikan_usd"] == b["total_penarikan_usd"]
    r2 = klien.post("/api/penarikan", json={"investor_id": b["id"], "tanggal": hari(), "jumlah_usd": 100,
                                            "kurs": 16000, "catatan": "uji"})
    assert r2.status_code == 200, r2.text
    b2 = investor(klien, "B")
    assert b2["total_penarikan_usd"] == b["total_penarikan_usd"] + 100
    assert b2["uang_diterima_idr"] == b["uang_diterima_idr"] + 1485480
    assert b2["saldo_usd"] < b["saldo_usd"]


def test_penarikan_melebihi_saldo_ditolak(klien):
    b = investor(klien, "B")
    r = klien.post("/api/penarikan", json={"investor_id": b["id"], "tanggal": hari(), "jumlah_usd": b["saldo_usd"] + 1000,
                                           "kurs": 16000})
    assert r.status_code == 400 and "melebihi saldo" in r.json()["detail"]
    p = klien.post("/api/penarikan/pratinjau", json={"investor_id": b["id"], "jumlah_usd": b["saldo_usd"] + 1000}).json()
    assert p["melebihi_saldo"] and not p["bisa_disimpan"]
    assert "melebihi" in klien.post("/api/penarikan", json={"investor_id": b["id"], "tanggal": hari(),
            "jumlah_usd": 99999, "kurs": 16000}).json()["detail"]


def test_validasi_penarikan(klien):
    b = investor(klien, "B")
    ok = {"investor_id": b["id"], "tanggal": hari(), "jumlah_usd": 10, "kurs": 16000}
    assert "investor" in klien.post("/api/penarikan", json={**ok, "investor_id": ""}).json()["detail"].lower()
    assert "Jumlah" in klien.post("/api/penarikan", json={**ok, "jumlah_usd": 0}).json()["detail"]
    assert "Kurs" in klien.post("/api/penarikan", json={**ok, "kurs": 0}).json()["detail"]
    assert "Tanggal" in klien.post("/api/penarikan", json={**ok, "tanggal": "2026-13-45"}).json()["detail"]
    assert "terlalu kecil" in klien.post("/api/penarikan", json={**ok, "jumlah_usd": 0.5}).json()["detail"]


def test_penghapusan_modal_yang_menyebabkan_saldo_negatif_ditolak(klien):
    a = investor(klien, "A")
    trx = klien.get(f"/api/investor/{a['id']}").json()["transaksi"]
    tarik = next(t for t in trx if t["jenis"] == "penarikan")
    # tambah modal besar lalu tarik hampir semuanya, lalu coba hapus modal tambahan
    m = klien.post("/api/transaksi/modal", json={"investor_id": a["id"], "jenis": "tambahan", "tanggal": hari(-2),
                                                 "jumlah_usd": 5000, "kurs": 16000}).json()
    tersedia = klien.post("/api/penarikan/pratinjau", json={"investor_id": a["id"], "jumlah_usd": 1}).json()["saldo_tersedia_usd"]
    assert klien.post("/api/penarikan", json={"investor_id": a["id"], "tanggal": hari(), "jumlah_usd": tersedia}).status_code == 200
    r = klien.delete(f"/api/transaksi/{m['id']}")
    assert r.status_code == 400 and "negatif" in r.json()["detail"]
    # modal awal tidak dapat dihapus
    awal = next(t for t in klien.get(f"/api/investor/{a['id']}").json()["transaksi"] if t["jenis"] == "modal_awal")
    assert klien.delete(f"/api/transaksi/{awal['id']}").status_code == 400
    # penarikan boleh dihapus
    assert klien.delete(f"/api/transaksi/{tarik['id']}").status_code == 200


def test_pengaturan_tersimpan_dan_dipakai(klien):
    s = klien.get("/api/pengaturan").json()
    assert s["persen"]["senin"] == 4 and s["persen"]["sabtu"] == 2 and s["persen"]["minggu"] == 0
    assert s["biaya_penarikan_persen"] == 6 and s["biaya_konversi_persen"] == 0.5 and s["biaya_transfer_idr"] == 11000
    s["persen"]["senin"] = 5
    s["biaya_transfer_idr"] = 20000
    s["kurs"] = 15500
    r = klien.put("/api/pengaturan", json=s)
    assert r.status_code == 200, r.text
    s2 = klien.get("/api/pengaturan").json()
    assert s2["persen"]["senin"] == 5 and s2["biaya_transfer_idr"] == 20000 and s2["kurs"] == 15500
    b = investor(klien, "B")
    p = klien.post("/api/penarikan/pratinjau", json={"investor_id": b["id"], "jumlah_usd": 100}).json()["rincian"]
    assert p["biaya_transfer_idr"] == 20000 and p["kurs"] == 15500
    assert klien.get("/api/pemeriksaan").json()["ok"]


def test_validasi_pengaturan(klien):
    s = klien.get("/api/pengaturan").json()
    for ubah, kata in [(lambda x: x["persen"].__setitem__("senin", 150), "Senin"),
                       (lambda x: x["persen"].__setitem__("selasa", "abc"), "Selasa"),
                       (lambda x: x.__setitem__("biaya_penarikan_persen", -1), "Biaya Penarikan"),
                       (lambda x: x.__setitem__("kurs", 0), "Kurs")]:
        x = json.loads(json.dumps(s))
        ubah(x)
        r = klien.put("/api/pengaturan", json=x)
        assert r.status_code == 400 and kata in r.json()["detail"], r.text


def test_kurs_manual_dan_pembaruan_online_gagal_dengan_sopan(klien, monkeypatch):
    assert klien.post("/api/kurs", json={"nilai": 16250}).json()["kurs"] == 16250
    assert klien.post("/api/kurs", json={"nilai": 0}).status_code == 400
    import urllib.request

    def gagal(*a, **k):
        raise OSError("tidak ada internet")
    monkeypatch.setattr(urllib.request, "urlopen", gagal)
    r = klien.post("/api/kurs/perbarui")
    assert r.status_code == 200 and r.json()["ok"] is False
    assert klien.get("/api/pengaturan").json()["kurs"] == 16250
    assert len(klien.get("/api/kurs/riwayat").json()["riwayat"]) >= 2


def test_hasil_aktual_dan_perbandingan(klien):
    a = investor(klien, "A")
    tgl = hari(-3)
    assert klien.post("/api/aktual", json={"tanggal": tgl, "persen": -2, "catatan": "rugi"}).status_code == 200
    assert klien.post("/api/aktual", json={"tanggal": tgl, "persen": 500}).status_code == 400
    banding = klien.get(f"/api/aktual/perbandingan?investor_id={a['id']}").json()
    assert banding["ringkas"]["selisih_usd"] < 0
    assert banding["ringkas"]["selisih_persen"] < 0
    assert banding["ringkas"]["selisih_idr"] < 0
    h = klien.get(f"/api/harian/investor/{a['id']}").json()["baris"]
    baris = next(b for b in h if b["tanggal"] == tgl)
    assert baris["persen"] == -2 and baris["sumber_persen"] == "aktual"
    assert klien.delete(f"/api/aktual/{tgl}").status_code == 200
    assert klien.get(f"/api/aktual/perbandingan?investor_id={a['id']}").json()["ringkas"]["selisih_usd"] == 0


def test_perhitungan_harian_per_tanggal(klien):
    d = klien.get(f"/api/harian?tanggal={hari()}").json()
    assert len(d["baris"]) == 3
    for b in d["baris"]:
        assert abs(b["saldo_awal"] + b["tambahan"] + b["penyesuaian"] - b["penarikan"] + b["keuntungan"]
                   - b["saldo_akhir"]) < 0.005
    assert klien.get("/api/harian?tanggal=salah").status_code == 400


def test_riwayat_transaksi_filter_dan_cari(klien):
    semua = klien.get("/api/transaksi").json()["transaksi"]
    assert {t["jenis"] for t in semua} >= {"modal_awal", "tambahan", "penarikan"}
    a = investor(klien, "A")
    hanya_a = klien.get(f"/api/transaksi?investor_id={a['id']}").json()["transaksi"]
    assert hanya_a and all(t["investor_id"] == a["id"] for t in hanya_a)
    assert all(t["jenis"] == "penarikan" for t in klien.get("/api/transaksi?jenis=penarikan").json()["transaksi"])
    assert len(klien.get("/api/transaksi?q=contoh penarikan").json()["transaksi"]) == 1


def test_simulator_dan_skenario(klien):
    s = klien.get("/api/pengaturan").json()
    body = {"modal_awal": 1000, "tanggal_mulai": "2026-10-05", "tanggal_akhir": "2026-10-11", "kurs": 16000,
            "persen": {"senin": 4, "selasa": 4, "rabu": 4, "kamis": 4, "jumat": 4, "sabtu": 2, "minggu": 0}}
    r = klien.post("/api/simulator", json=body).json()
    assert r["baris"][0]["saldo_akhir"] == 1040 and r["baris"][1]["saldo_akhir"] == 1081.6
    assert r["saldo_akhir"] == r["baris"][-1]["saldo_akhir"]
    assert abs(r["persen_keuntungan"] - (r["saldo_akhir"] - 1000) / 10) < 1e-9
    assert klien.post("/api/simulator", json={**body, "tanggal_akhir": "2026-10-01"}).status_code == 400
    assert klien.post("/api/simulator", json={**body, "modal_awal": -1}).status_code == 400
    sk = klien.post("/api/skenario", json={"modal_awal": 1000, "tanggal_mulai": "2026-10-05", "skenario": [
        {"nama": "KONSERVATIF", "senin_jumat": 1, "sabtu": 0.5, "minggu": 0},
        {"nama": "DASAR", "senin_jumat": 4, "sabtu": 2, "minggu": 0},
        {"nama": "OPTIMISTIS", "senin_jumat": 5, "sabtu": 3, "minggu": 0}]}).json()
    nilai = [x["horizon"][0]["saldo"] for x in sk["skenario"]]
    assert nilai[0] < nilai[1] < nilai[2]
    assert [h["hari"] for h in sk["skenario"][0]["horizon"]] == [30, 60, 90, 180, 365]


def test_grafik(klien):
    g = klien.get("/api/grafik").json()
    assert len(g["investor"]) == 3 and all(len(i["saldo_usd"]) == len(g["tanggal"]) for i in g["investor"])
    g2 = klien.get(f"/api/grafik?investor_ids={g['investor'][0]['id']}&status=aktif").json()
    assert len(g2["investor"]) == 1


def test_cadangan_ekspor_impor(klien):
    json_ = klien.get("/api/cadangan/json")
    assert json_.status_code == 200
    isi = json_.json()
    assert len(isi["tabel"]["investor"]) == 3
    sebelum = klien.get("/api/dashboard").json()["total"]
    # ubah data, lalu pulihkan
    a = investor(klien, "A")
    klien.delete(f"/api/investor/{a['id']}")
    assert klien.get("/api/dashboard").json()["total"]["jumlah_investor"] == 2
    r = klien.post("/api/cadangan/impor", files={"berkas": ("c.json", json.dumps(isi), "application/json")})
    assert r.status_code == 200, r.text
    setelah = klien.get("/api/dashboard").json()["total"]
    assert setelah == sebelum
    assert klien.get("/api/pemeriksaan").json()["ok"]
    # file salah ditolak tanpa mengubah data
    bad = klien.post("/api/cadangan/impor", files={"berkas": ("c.json", b"bukan json", "application/json")})
    assert bad.status_code == 400
    bad = klien.post("/api/cadangan/impor", files={"berkas": ("c.json", json.dumps({"a": 1}), "application/json")})
    assert bad.status_code == 400
    assert klien.get("/api/dashboard").json()["total"] == sebelum


def test_ekspor_csv_dan_excel(klien):
    for t in ("investor", "transaksi", "harian"):
        r = klien.get(f"/api/cadangan/csv?tabel={t}")
        assert r.status_code == 200 and len(r.text.splitlines()) > 1
    r = klien.get("/api/cadangan/xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Investor", "Transaksi", "Perhitungan Harian", "Pengaturan"} <= set(wb.sheetnames)
    assert wb["Investor"].max_row == 4
    assert klien.get("/api/cadangan/csv?tabel=aneh").status_code == 400


def test_ganti_kata_sandi(klien):
    assert klien.post("/api/ganti-sandi", json={"sandi_lama": "salah", "sandi_baru": "sandibaru123"}).status_code == 400
    assert klien.post("/api/ganti-sandi", json={"sandi_lama": "admin123", "sandi_baru": "pendek"}).status_code == 400
    assert klien.post("/api/ganti-sandi", json={"sandi_lama": "admin123", "sandi_baru": "sandibaru123"}).status_code == 200
    assert klien.get("/api/saya").json()["sandi_awal"] is False


def test_pemeriksaan_mendeteksi_data_rusak(klien):
    from app.db import sambung
    conn = sambung()
    conn.execute("UPDATE perhitungan_harian SET saldo_akhir = '1' WHERE rowid = (SELECT MIN(rowid) FROM perhitungan_harian)")
    conn.commit()
    r = klien.get("/api/pemeriksaan").json()
    assert r["ok"] is False and r["jumlah_masalah"] >= 1
